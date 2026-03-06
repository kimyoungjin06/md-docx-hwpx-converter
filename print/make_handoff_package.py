#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Optional, Tuple

import docx
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

from lxml import etree


DOCX_XML_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "v": "urn:schemas-microsoft-com:vml",
}


CAPTION_RE = re.compile(r"^\s*그림\s*([0-9]+(?:[-\.][0-9]+)*|[A-Za-z]+(?:[-\.][0-9]+)+)\s*[\.:]?\s*(.*)$")
RID_RE = re.compile(r"\brId\d+\.[A-Za-z0-9]+\b")


def _run(cmd: List[str], *, cwd: Optional[Path] = None) -> None:
    subprocess.run(cmd, check=True, cwd=str(cwd) if cwd else None)


def _iter_block_items(parent: DocxDocument | _Cell) -> Iterator[Paragraph | Table]:
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _iter_paragraphs(parent: DocxDocument | _Cell) -> Iterator[Paragraph]:
    for block in _iter_block_items(parent):
        if isinstance(block, Paragraph):
            yield block
            continue
        if isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    yield from _iter_paragraphs(cell)


def _paragraph_image_rids(paragraph: Paragraph) -> List[str]:
    rids: List[str] = []
    try:
        # python-docx's BaseOxmlElement.xpath() uses its own namespace map; do not pass namespaces=...
        rids.extend(paragraph._p.xpath(".//a:blip/@r:embed"))
        # Fallback for legacy VML images (prefix not registered in python-docx nsmap).
        rids.extend(paragraph._p.xpath('.//*[local-name()="imagedata"]/@r:id'))
    except Exception:
        return []
    return [rid for rid in rids if isinstance(rid, str) and rid]


@dataclass
class ParaInfo:
    text: str
    image_rids: List[str]


@dataclass
class ImageEntry:
    img_idx: int
    new_filename: str
    docx_partname: str
    content_type: str
    occurrence_count: int
    caption_guess: str
    figure_no: str
    first_para_text: str


def _guess_caption(paras: List[ParaInfo], idx: int) -> str:
    def is_caption(text: str) -> bool:
        return bool(CAPTION_RE.match(text))

    here = paras[idx].text.strip()
    if here and is_caption(here):
        return here

    if idx - 1 >= 0:
        prev = paras[idx - 1].text.strip()
        if prev and is_caption(prev):
            return prev

    if idx + 1 < len(paras):
        nxt = paras[idx + 1].text.strip()
        if nxt and is_caption(nxt):
            return nxt

    return ""


def _figure_no_from_caption(caption: str) -> str:
    m = CAPTION_RE.match(caption.strip())
    if not m:
        return ""
    fig = (m.group(1) or "").strip()
    fig = fig.replace(".", "-").upper()
    fig = re.sub(r"[^0-9A-Z\-]+", "", fig)
    return fig


def _image_filename(*, idx: int, ext: str, figure_no: str) -> str:
    ext = ext.lower()
    if not ext.startswith("."):
        ext = f".{ext}" if ext else ""

    if figure_no:
        return f"FIG_{idx:04d}_{figure_no}{ext}"
    return f"FIG_{idx:04d}{ext}"


def extract_images_in_order(*, docx_path: Path, out_dir: Path) -> List[ImageEntry]:
    document = docx.Document(str(docx_path))
    paras: List[ParaInfo] = []

    for p in _iter_paragraphs(document):
        text = (p.text or "").strip()
        rids = _paragraph_image_rids(p)
        paras.append(ParaInfo(text=text, image_rids=rids))

    # Keep first-appearance order of unique image parts.
    part_by_key: Dict[str, object] = {}
    first_para_text_by_key: Dict[str, str] = {}
    caption_by_key: Dict[str, str] = {}
    occ_count_by_key: Dict[str, int] = {}
    order: List[str] = []

    for idx, para in enumerate(paras):
        if not para.image_rids:
            continue
        caption_guess = _guess_caption(paras, idx)
        for rid in para.image_rids:
            part = document.part.related_parts.get(rid)
            if part is None:
                continue
            content_type = getattr(part, "content_type", "")
            if not isinstance(content_type, str) or not content_type.startswith("image/"):
                continue
            partname = getattr(part, "partname", "")
            key = str(partname) if partname else str(rid)
            if key not in part_by_key:
                part_by_key[key] = part
                first_para_text_by_key[key] = para.text
                caption_by_key[key] = caption_guess
                occ_count_by_key[key] = 1
                order.append(key)
            else:
                occ_count_by_key[key] = occ_count_by_key.get(key, 0) + 1
                if not caption_by_key.get(key) and caption_guess:
                    caption_by_key[key] = caption_guess

    out_dir.mkdir(parents=True, exist_ok=True)
    entries: List[ImageEntry] = []

    for img_idx, key in enumerate(order, start=1):
        part = part_by_key[key]
        partname = getattr(part, "partname", "")
        content_type = getattr(part, "content_type", "")
        blob = getattr(part, "blob", b"")
        ext = Path(str(partname)).suffix
        if not ext:
            # Fallback to content-type
            ext = {
                "image/png": ".png",
                "image/jpeg": ".jpg",
                "image/jpg": ".jpg",
                "image/gif": ".gif",
                "image/bmp": ".bmp",
                "image/tiff": ".tif",
                "image/x-emf": ".emf",
                "image/x-wmf": ".wmf",
            }.get(str(content_type).lower(), "")

        caption_guess = caption_by_key.get(key, "")
        figure_no = _figure_no_from_caption(caption_guess)
        filename = _image_filename(idx=img_idx, ext=ext, figure_no=figure_no)

        out_path = out_dir / filename
        out_path.write_bytes(blob if isinstance(blob, (bytes, bytearray)) else b"")

        entries.append(
            ImageEntry(
                img_idx=img_idx,
                new_filename=filename,
                docx_partname=str(partname),
                content_type=str(content_type),
                occurrence_count=occ_count_by_key.get(key, 1),
                caption_guess=caption_guess,
                figure_no=figure_no,
                first_para_text=first_para_text_by_key.get(key, ""),
            )
        )

    return entries


def _rewrite_hwpx_image_placeholders(hwpx_path: Path, rid_to_filename: Dict[str, str]) -> None:
    if not rid_to_filename:
        return

    rid_to_filename_lc = {k.lower(): v for k, v in rid_to_filename.items()}

    with zipfile.ZipFile(hwpx_path) as zf_in:
        section_xml = zf_in.read("Contents/section0.xml")
        root = etree.fromstring(section_xml)  # noqa: S320 (trusted local file)

        changed = 0
        for t_el in root.xpath('//*[local-name()="t"]'):
            text = t_el.text
            if not text:
                continue

            def repl(match: re.Match[str]) -> str:
                nonlocal changed
                key = match.group(0)
                replacement = rid_to_filename_lc.get(key.lower())
                if replacement:
                    changed += 1
                    return replacement
                return key

            new_text = RID_RE.sub(repl, text)
            if new_text != text:
                t_el.text = new_text

        if changed == 0:
            return

        new_section_xml = etree.tostring(
            root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )

    tmp_path = hwpx_path.with_suffix(".tmp.hwpx")
    with zipfile.ZipFile(hwpx_path) as zf_in, zipfile.ZipFile(tmp_path, "w") as zf_out:
        for info in zf_in.infolist():
            data = zf_in.read(info.filename)
            if info.filename == "Contents/section0.xml":
                data = new_section_xml
            zf_out.writestr(info, data)
    tmp_path.replace(hwpx_path)


def make_handoff_package(
    *,
    docx_path: Path,
    output_dir: Path,
    template_hwpx: Optional[Path],
    hwpx_mode: str,
    include_pdf: bool,
    include_hwpx: bool,
    include_images: bool,
    base_name: Optional[str],
) -> Path:
    docx_path = docx_path.resolve()
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    stem = base_name.strip() if base_name else docx_path.stem
    stem = re.sub(r"[\\/:*?\"<>|]+", "_", stem).strip()
    if not stem:
        stem = "report"

    doc_dir = output_dir / "document"
    img_dir = output_dir / "images"
    map_dir = output_dir / "mappings"
    doc_dir.mkdir(parents=True, exist_ok=True)
    map_dir.mkdir(parents=True, exist_ok=True)

    out_docx = doc_dir / f"{stem}.docx"
    shutil.copy2(docx_path, out_docx)

    image_entries: List[ImageEntry] = []
    if include_images:
        if img_dir.exists():
            shutil.rmtree(img_dir)
        image_entries = extract_images_in_order(docx_path=out_docx, out_dir=img_dir)
        manifest = map_dir / "images_manifest.csv"
        with manifest.open("w", newline="", encoding="utf-8") as fp:
            writer = csv.DictWriter(
                fp,
                fieldnames=[
                    "img_idx",
                    "new_filename",
                    "figure_no",
                    "caption_guess",
                    "occurrence_count",
                    "docx_partname",
                    "content_type",
                    "first_para_text",
                ],
            )
            writer.writeheader()
            for e in image_entries:
                writer.writerow(
                    {
                        "img_idx": e.img_idx,
                        "new_filename": e.new_filename,
                        "figure_no": e.figure_no,
                        "caption_guess": e.caption_guess,
                        "occurrence_count": e.occurrence_count,
                        "docx_partname": e.docx_partname,
                        "content_type": e.content_type,
                        "first_para_text": e.first_para_text,
                    }
                )

    if include_hwpx:
        if not template_hwpx:
            raise ValueError("--template-hwpx is required to generate HWPX.")
        out_hwpx = doc_dir / f"{stem}.hwpx"
        minifier = Path(__file__).resolve().parent / "make_hwpx_template.py"
        converter = Path(__file__).resolve().parent / "convert_docx_to_hwpx.py"
        with tempfile.TemporaryDirectory(prefix="hwpx_template_") as tmpdir:
            slim_template = Path(tmpdir) / "template.hwpx"
            _run([sys.executable, str(minifier), str(template_hwpx.resolve()), "--output", str(slim_template)])
            _run(
                [
                    sys.executable,
                    str(converter),
                    str(out_docx),
                    "--output",
                    str(out_hwpx),
                    "--template-hwpx",
                    str(slim_template),
                    "--mode",
                    hwpx_mode,
                ]
            )

        if image_entries:
            rid_to_filename = {Path(e.docx_partname).name: e.new_filename for e in image_entries if e.docx_partname}
            _rewrite_hwpx_image_placeholders(out_hwpx, rid_to_filename)

    if include_pdf:
        out_pdf = doc_dir / f"{stem}.pdf"
        pdf_script = Path(__file__).resolve().parent / "docx_to_review_pdf.sh"
        _run(["bash", str(pdf_script), str(out_docx), str(out_pdf)])

    readme = output_dir / "00_README.md"
    lines = [
        "# 편집사 전달 패키지",
        "",
        "- `document/`: 편집본(DOCX/HWPX) + 검토용 PDF",
        "- `images/`: DOCX에서 추출한 그림(문서 등장 순서 기준으로 리네이밍)",
        "- `mappings/images_manifest.csv`: 그림 파일명 ↔ 캡션/도큐먼트 내부 매핑",
        "",
        "## 주의사항",
        "",
        "- `*.hwpx`는 best-effort 변환 결과이며, 표/그림/수식 레이아웃은 편집 과정에서 조정이 필요할 수 있다.",
        "- `document/*.hwpx`에는 그림이 placeholder로 들어가 있으므로, `images/FIG_0001_...` 순서대로 재삽입하는 것을 권장한다.",
        "- 편집사가 `.hwp`만 받을 경우: 한글(Hancom)에서 `document/*.hwpx`를 열고 `.hwp`로 저장해서 전달한다.",
        "- `images/`는 재삽입/재편집을 위한 원본 추출본이다(문서 내 포함 이미지와 동일).",
        "",
    ]
    if image_entries:
        lines.append(f"- 총 추출 이미지 수: {len(image_entries)}")
        lines.append("")
    readme.write_text("\n".join(lines), encoding="utf-8")

    return output_dir


def main() -> int:
    parser = argparse.ArgumentParser(description="Create a zipped-ready handoff folder for HWP-based editors.")
    parser.add_argument("--docx", type=Path, required=True, help="Input .docx path")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Output directory (default: dist/handoff_<YYYYMMDD>_<docx_stem>).",
    )
    parser.add_argument("--template-hwpx", type=Path, default=None, help="Template .hwpx used for DOCX→HWPX.")
    parser.add_argument(
        "--hwpx-mode",
        choices=["outline", "report"],
        default="outline",
        help="DOCX→HWPX heading style mapping (default: outline).",
    )
    parser.add_argument("--no-pdf", action="store_true", help="Do not generate review PDF.")
    parser.add_argument("--no-hwpx", action="store_true", help="Do not generate HWPX.")
    parser.add_argument("--no-images", action="store_true", help="Do not extract/rename images from DOCX.")
    parser.add_argument("--base-name", type=str, default=None, help="Base filename for document outputs.")

    args = parser.parse_args()

    today = datetime.now().strftime("%Y%m%d")
    default_dir = Path("dist") / f"handoff_{today}_{args.docx.stem}"
    output_dir = args.output_dir or default_dir

    out = make_handoff_package(
        docx_path=args.docx,
        output_dir=output_dir,
        template_hwpx=args.template_hwpx,
        hwpx_mode=args.hwpx_mode,
        include_pdf=not args.no_pdf,
        include_hwpx=not args.no_hwpx,
        include_images=not args.no_images,
        base_name=args.base_name,
    )
    print(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
